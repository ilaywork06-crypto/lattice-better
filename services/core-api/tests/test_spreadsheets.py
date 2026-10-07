import io

from openpyxl import load_workbook


def _workbook(admin, query=""):
    r = admin.get(f"/spreadsheets/import-template{query}")
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def _upload(api, wb, name="items.xlsx"):
    buf = io.BytesIO()
    wb.save(buf)
    return api.post("/spreadsheets/import", files={"file": (name, buf.getvalue())})


def test_import_headers_are_the_templates_creation_fields(admin):
    wb = _workbook(admin)
    assert "C-PRB Power Regulator Board" in wb.sheetnames
    ws = wb["C-PRB Power Regulator Board"]
    assert [c.value for c in ws[1]] == ["Serial", "Project", "Revision", "Location"]
    ws = wb["A-SPM Signal Processing Module"]
    assert [c.value for c in ws[1]][-1] == "Contents"


def test_import_creates_items_and_their_contents(admin):
    wb = _workbook(admin)
    wb["C-PRB Power Regulator Board"].append([None, "Sparrow", "03", "Lab A"])
    wb["C-PRB Power Regulator Board"].append(["C-PRB-050", None, None, None])
    wb["A-SPM Signal Processing Module"].append([None, "Lab A", "C-PRB-001, C-PRB-050"])
    r = _upload(admin, wb)
    assert r.status_code == 200, r.text
    assert r.json() == {"created": 3, "by_template": {"Power Regulator Board": 2,
                                                      "Signal Processing Module": 1}}
    module = admin.ok("GET", "/items?type=assembly")["items"][0]
    assert module["children_count"] == 2 and module["missing_children"] == 0


def test_import_is_all_or_nothing_and_names_every_bad_cell(admin):
    wb = _workbook(admin)
    ws = wb["C-PRB Power Regulator Board"]
    ws.append([None, "Sparrow", "03", "Lab A"])
    ws.append([None, "Nope", "abc", "Lab A"])
    ws.append(["X-1", None, None, None])
    r = _upload(admin, wb)
    assert r.status_code == 400
    cells = {(i["cell"], i["column"]) for i in r.json()["error"]["issues"]}
    assert {("B3", "Project"), ("C3", "Revision"), ("A4", "Serial")} <= cells
    assert admin.ok("GET", "/items")["total"] == 0


def test_import_rejects_unknown_and_missing_columns(admin):
    wb = _workbook(admin, f"?template_id={admin.template('Resistor Pack')['id']}")
    ws = wb["C-RES Resistor Pack"]
    ws.delete_cols(2)
    ws["B1"] = "Colour"
    ws.append([None, "red"])
    messages = [i["message"] for i in _upload(admin, wb).json()["error"]["issues"]]
    assert any("not a field" in m for m in messages)
    assert any("required column 'Quantity'" in m for m in messages)


def test_import_accepts_only_excel_and_only_managers(world, admin):
    r = admin.post("/spreadsheets/import", files={"file": ("x.csv", b"a,b")})
    assert r.status_code == 400
    wb = _workbook(admin)
    assert _upload(world.as_("editor"), wb).status_code == 403


def test_export_has_a_sheet_per_template_with_readable_values(admin):
    admin.create_item("Power Regulator Board", values={"revision": "09"})
    r = admin.get("/spreadsheets/export?type=card")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["C-PRB Power Regulator Board"]
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]], strict=True))
    assert (row["Serial"], row["Project"], row["Managers"], row["Revision"]) == (
        "C-PRB-001", "Falcon", "Noa Manager", "RV-09")
    assert "attachment" in r.headers["content-disposition"]
