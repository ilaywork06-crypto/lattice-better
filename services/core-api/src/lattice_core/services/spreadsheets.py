"""Excel import and export (openpyxl).

**Import workbook:** one sheet per template; the header row is exactly that
template's creation fields (per-item and list fields, in order), preceded by an
optional *Serial* column and followed, for containers, by *Contents* (serials of
items to place inside). Required headers are highlighted, every header carries
a comment with its format and allowed values, and list columns get a dropdown.

**Import is all-or-nothing:** every row goes through the same use case the UI
uses; if any cell is wrong nothing is saved, and the error lists *every* bad
cell (sheet, A1 reference, reason).
"""

from __future__ import annotations

import io
import re
from datetime import UTC
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from lattice_core.db.models import Item, ItemTemplate, TemplateField
from lattice_core.domain.enums import (
    CATALOG_FIELD_TYPES,
    SERIAL_LETTER,
    FieldMode,
    FieldType,
    ItemState,
    ItemType,
)
from lattice_core.domain.errors import Issue, LatticeError, RuleViolation, ValidationFailed
from lattice_core.domain.fields import LETTERS
from lattice_core.domain.permissions import Permission
from lattice_core.repositories.audit import AuditFilter
from lattice_core.schemas.items import ItemCreate
from lattice_core.schemas.misc import ImportResult
from lattice_core.services.base import Service
from lattice_core.services.item_values import effective_value

META_SHEET = "_lattice"
README_SHEET = "Read me"
SERIAL_HEADER = "Serial"
CONTENTS_HEADER = "Contents"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_TYPE_ORDER = {ItemType.CARD: 0, ItemType.ASSEMBLY: 1, ItemType.SETUP: 2}
_BAD_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
_REQUIRED_FILL = PatternFill("solid", fgColor="FFE9E3F8")
_HEADER_FILL = PatternFill("solid", fgColor="FFEFF1F5")
_PHYSICAL = (FieldType.STATUS, FieldType.LOCATION, FieldType.PARENT)


def sheet_title(tpl: ItemTemplate) -> str:
    """Unique per template ((type, prefix) is unique) and within Excel's rules."""
    head = f"{SERIAL_LETTER[tpl.type]}-{tpl.serial_prefix} "
    return (head + _BAD_TITLE_CHARS.sub(" ", tpl.name))[:31].strip()


def creation_fields(tpl: ItemTemplate) -> list[TemplateField]:
    """Fields filled in when an item is created (files can't ride in a cell)."""
    return [f for f in tpl.fields if f.mode != FieldMode.FIXED and f.field_type != FieldType.FILES]


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _is_blank(row) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


class SpreadsheetService(Service):
    def _templates(self, template_id: int | None, item_type: ItemType | None):
        if template_id is not None:
            templates = [self.uow.templates.require(template_id)]
        else:
            templates = self.uow.templates.list(type=item_type)
        return sorted(templates, key=lambda t: (_TYPE_ORDER[t.type], t.name.lower()))

    def filename(self, kind: str, template_id: int | None, item_type: ItemType | None) -> str:
        if template_id is not None:
            tpl = self.uow.templates.require(template_id)
            return f"lattice_{kind}_{tpl.type.value}_{tpl.serial_prefix}.xlsx"
        return f"lattice_{kind}{f'_{item_type.value}' if item_type else ''}.xlsx"

    # ─────────────────────────── import workbook ───────────────────────────
    def _allowed(self, f: TemplateField) -> list[str]:
        if f.field_type == FieldType.ENUM or (f.mode == FieldMode.CHOICE and f.options):
            return self.services.values.options_display(f.field_type, f.mode, f.options)
        if f.field_type == FieldType.STATUS:
            return [s.value for s in ItemState]
        if f.field_type == FieldType.LETTER:
            return list(LETTERS)
        if f.field_type == FieldType.BOOLEAN:
            return ["yes", "no"]
        if f.field_type in CATALOG_FIELD_TYPES:
            return [o.value for o in self.uow.catalog.list(CATALOG_FIELD_TYPES[f.field_type],
                                                           active_only=True)]
        if f.field_type == FieldType.LOCATION:
            return [loc.name for loc in self.uow.locations.list()]
        return []

    def _hint(self, f: TemplateField) -> str:
        parts = [f"Type: {f.field_type.value}", "Required" if f.required else "Optional"]
        config = f.config or {}
        if config.get("pattern"):
            parts.append(f"Format {config['pattern']}: type only the digits (#)")
        hints = {
            FieldType.DESCRIPTION: f"At least {config.get('min_length', 8)} non-blank characters",
            FieldType.MANAGERS: "Manager emails or names, separated by commas",
            FieldType.RESPONSIBLE: "A user's email or name",
            FieldType.PARENT: "The serial of the item this one goes into",
            FieldType.DATE: "YYYY-MM-DD",
        }
        if f.field_type in hints:
            parts.append(hints[f.field_type])
        if f.mode == FieldMode.CHOICE:
            parts.append("Blank = the first value of the list")
        allowed = self._allowed(f)
        if allowed:
            more = " …" if len(allowed) > 25 else ""
            parts.append("Allowed: " + ", ".join(allowed[:25]) + more)
        return "\n".join(parts)

    def import_workbook(self, template_id: int | None = None,
                        item_type: ItemType | None = None) -> bytes:
        self.services.require(Permission.READ)
        templates = self._templates(template_id, item_type)
        if not templates:
            raise RuleViolation("There are no templates to build an import file from")
        wb = Workbook()
        readme = wb.active
        readme.title = README_SHEET
        readme.append(["Lattice import"])
        readme["A1"].font = Font(bold=True, size=14)
        for line in (
            "One sheet per template. Each row creates one item from that template.",
            "The header row lists the fields filled in when an item is created; highlighted "
            "headers are required. Hover a header to see its format.",
            f"'{SERIAL_HEADER}': leave blank to get the next serial automatically.",
            f"'{CONTENTS_HEADER}': serials of existing items (or of items from earlier rows or "
            "sheets) to place inside, separated by commas.",
            "Fields set on the template itself are not listed — every item gets them.",
            "If any cell is invalid nothing is imported, and every invalid cell is listed.",
        ):
            readme.append([line])
        readme.column_dimensions["A"].width = 110

        meta = wb.create_sheet(META_SHEET)
        meta.append(["sheet", "template_id"])
        meta.sheet_state = "hidden"

        for tpl in templates:
            ws = wb.create_sheet(sheet_title(tpl))
            meta.append([ws.title, tpl.id])
            fields = creation_fields(tpl)
            headers = [(SERIAL_HEADER, "Optional. Blank = next serial, e.g. "
                        f"{self.services.templates.next_serial(tpl)}")]
            headers += [(f.label, self._hint(f)) for f in fields]
            if tpl.child_links:
                names = ", ".join(_slot_hint(link) for link in tpl.child_links)
                headers.append((CONTENTS_HEADER, f"Serials of items to place inside ({names})"))
            ws.append([h for h, _ in headers])
            for col, (_, hint) in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col)
                cell.font = Font(bold=True)
                cell.fill = _HEADER_FILL
                cell.comment = Comment(hint, "Lattice")
                width = max(14, len(str(cell.value)) + 4)
                ws.column_dimensions[get_column_letter(col)].width = width
            for col, f in enumerate(fields, start=2):
                if f.required:
                    ws.cell(row=1, column=col).fill = _REQUIRED_FILL
                allowed = self._allowed(f)
                formula = '"' + ",".join(a.replace('"', "") for a in allowed) + '"'
                if allowed and len(formula) <= 255 and f.field_type != FieldType.MANAGERS:
                    dv = DataValidation(type="list", formula1=formula, allow_blank=not f.required)
                    ws.add_data_validation(dv)
                    letter = get_column_letter(col)
                    dv.add(f"{letter}2:{letter}1000")
            ws.freeze_panes = "A2"
        return _save(wb)

    # ─────────────────────────── import ───────────────────────────
    def _resolve_sheets(self, wb) -> list[tuple[Any, ItemTemplate | None]]:
        mapping: dict[str, int] = {}
        if META_SHEET in wb.sheetnames:
            for row in wb[META_SHEET].iter_rows(min_row=2, values_only=True):
                if row and row[0] and row[1]:
                    try:
                        mapping[str(row[0])] = int(row[1])
                    except (TypeError, ValueError):
                        continue
        by_letter = {letter: t for t, letter in SERIAL_LETTER.items()}
        out = []
        for ws in wb.worksheets:
            if ws.title in (META_SHEET, README_SHEET):
                continue
            tpl = self.uow.templates.get(mapping.get(ws.title))
            if tpl is None:
                m = re.match(r"^([CAS])-([A-Za-z]{3})\b", ws.title)
                if m:
                    tpl = self.uow.templates.by_type_and_prefix(by_letter[m.group(1)],
                                                                m.group(2).upper())
            if tpl is None:
                tpl = self.uow.templates.by_name(ws.title)
            out.append((ws, tpl))
        # Cards first, then assemblies, then setups: contents exist before containers.
        out.sort(key=lambda p: _TYPE_ORDER[p[1].type] if p[1] else -1)
        return out

    def import_items(self, content: bytes, filename: str) -> ImportResult:  # noqa: C901
        self.services.require(Permission.IMPORT_DATA)
        if not filename.lower().endswith((".xlsx", ".xlsm")):
            raise RuleViolation("Please upload an Excel .xlsx file")
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure means "not a workbook"
            raise RuleViolation("This is not a readable Excel (.xlsx) file") from exc

        issues: list[Issue] = []
        created = 0
        by_template: dict[str, int] = {}

        def err(sheet: str, row: int | None, col: int | None, message: str,
                column: str | None = None) -> None:
            letter = get_column_letter(col) if col else None
            issues.append(Issue(message=message, sheet=sheet, row=row, column=column,
                                cell=f"{letter}{row}" if letter and row else None))

        with self.uow.transaction():
            sheets = self._resolve_sheets(wb)
            if not sheets:
                raise RuleViolation("The workbook has no template sheets to import")
            for ws, tpl in sheets:
                if tpl is None:
                    err(ws.title, None, None,
                        "This sheet doesn't match any template — download a fresh import file")
                    continue
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                fields = creation_fields(tpl)
                by_label = {f.label.strip().lower(): f for f in fields}
                by_key = {f.key.lower(): f for f in fields}
                columns: dict[int, TemplateField] = {}
                serial_col = contents_col = None
                header_ok = True
                for col, raw in enumerate(rows[0], start=1):
                    name = re.sub(r"\s*\*$", "", str(raw or "").strip()).lower()
                    if not name:
                        continue
                    if name == SERIAL_HEADER.lower():
                        serial_col = col
                    elif name == CONTENTS_HEADER.lower():
                        contents_col = col
                    elif name in by_label or name in by_key:
                        columns[col] = by_label.get(name) or by_key[name]
                    else:
                        err(ws.title, 1, col, f"'{raw}' is not a field of '{tpl.name}'", str(raw))
                        header_ok = False
                present = set(columns.values())
                for f in fields:
                    if f.required and f not in present:
                        err(ws.title, 1, None, f"The required column '{f.label}' is missing",
                            f.label)
                        header_ok = False
                if not header_ok:
                    continue
                col_of_key = {f.key: c for c, f in columns.items()}

                for r, row in enumerate(rows[1:], start=2):
                    if _is_blank(row):
                        continue

                    def cell(c, row=row):
                        return row[c - 1] if c and c - 1 < len(row) else None

                    values = {f.key: cell(c) for c, f in columns.items()
                              if cell(c) not in (None, "")}
                    row_ok = True
                    children: list[Item] = []
                    contents = cell(contents_col)
                    if contents not in (None, ""):
                        for s in re.split(r"[,;\n]", str(contents)):
                            if not s.strip():
                                continue
                            child = self.uow.items.by_serial(s)
                            if child is None:
                                err(ws.title, r, contents_col,
                                    f"no item has the serial '{s.strip().upper()}'",
                                    CONTENTS_HEADER)
                                row_ok = False
                            else:
                                children.append(child)

                    serial = cell(serial_col)
                    savepoint = self.uow.session.begin_nested()
                    try:
                        item = self.services.items.create(ItemCreate(
                            template_id=tpl.id, values=values,
                            serial=str(serial) if serial not in (None, "") else None,
                        ))
                    except LatticeError as exc:
                        savepoint.rollback()
                        if exc.issues:
                            for i in exc.issues:
                                err(ws.title, r, col_of_key.get(i.field), i.message,
                                    i.label or i.field)
                        else:
                            is_serial = "serial" in exc.message.lower() and serial_col
                            err(ws.title, r, serial_col if is_serial else None, exc.message,
                                SERIAL_HEADER if is_serial else None)
                        continue
                    for child in children:
                        try:
                            self.services.hierarchy.place(child, item)
                        except LatticeError as exc:
                            err(ws.title, r, contents_col, exc.message, CONTENTS_HEADER)
                            row_ok = False
                    if not row_ok:
                        savepoint.rollback()
                        continue
                    savepoint.commit()
                    created += 1
                    by_template[tpl.name] = by_template.get(tpl.name, 0) + 1

            if issues:
                raise ValidationFailed.from_issues(
                    issues,
                    f"Nothing was imported: {len(issues)} problem(s) found. Fix the listed cells "
                    "and import the file again.",
                )
            if created:
                self.audit.record("import", f"Imported {created} item(s) from '{filename}'",
                                  details={"by_template": by_template})
        return ImportResult(created=created, by_template=by_template)

    # ─────────────────────────── export ───────────────────────────
    def export_items(self, template_id: int | None = None,
                     item_type: ItemType | None = None) -> bytes:
        """Every item, one sheet per template, with readable names instead of ids."""
        self.services.require(Permission.READ)
        wb = Workbook()
        wb.remove(wb.active)
        for tpl in self._templates(template_id, item_type):
            ws = wb.create_sheet(sheet_title(tpl))
            own_fields = [f for f in tpl.fields if f.field_type not in _PHYSICAL]
            headers = [SERIAL_HEADER, "State", "Location", "Inside"]
            headers += [f.label for f in own_fields]
            if tpl.child_links:
                headers.append(CONTENTS_HEADER)
            headers += ["Created", "Updated"]
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = _HEADER_FILL
            for it in self.uow.items.of_template(tpl.id):
                row: list[Any] = [
                    it.serial, it.state.value,
                    it.location.name if it.location else None,
                    it.parent.serial if it.parent else None,
                ]
                for f in own_fields:
                    shown = self.services.values.display(f.field_type,
                                                         effective_value(self.uow, it, f))
                    row.append(_cell(shown))
                if tpl.child_links:
                    row.append(", ".join(c.serial for c in it.children) or None)
                row += [_naive(it.created_at), _naive(it.updated_at)]
                ws.append(row)
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 18
            ws.freeze_panes = "A2"
        if not wb.worksheets:
            wb.create_sheet("items").append(["No templates yet"])
        return _save(wb)

    def export_audit(self, f: AuditFilter) -> bytes:
        self.services.require(Permission.READ)
        wb = Workbook()
        ws = wb.active
        ws.title = "audit"
        ws.append(["When (UTC)", "User", "Action", "Subject", "Summary", "Details"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for e in self.uow.audit.every(f):
            ws.append([_naive(e.created_at), e.user_name, e.action, e.subject, e.summary,
                       ", ".join(f"{k}: {v}" for k, v in (e.details or {}).items())])
        for col, width in zip("ABCDEF", (20, 22, 22, 32, 70, 60), strict=True):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        return _save(wb)


def _slot_hint(link) -> str:
    lo, hi = link.min_count, link.max_count
    if hi is None:
        return link.child.name + (f" — at least {lo}" if lo else "")
    return f"{link.child.name} — " + (f"exactly {lo}" if lo == hi else f"{lo} to {hi}")


def _cell(shown: Any) -> Any:
    if shown is None:
        return None
    if isinstance(shown, list):
        return ", ".join(str(d.get("name") if isinstance(d, dict) else d) for d in shown)
    if isinstance(shown, bool):
        return "yes" if shown else "no"
    return shown


def _naive(dt):
    if dt is None:
        return None
    return dt.astimezone(UTC).replace(tzinfo=None) if dt.tzinfo else dt
